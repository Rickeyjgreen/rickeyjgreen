from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


BARCODE_ALIASES = ("barcode (1)", "barcode", "copyright", "copyright code", "code")


class RosterError(ValueError):
    pass


@dataclass(frozen=True)
class RosterCodes:
    codes: frozenset[str]
    source_rows: int
    column: str


def normalize_code(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def _choose_header(rows: Iterable[Iterable[object]]) -> tuple[int, list[str], list[list[object]]]:
    materialized = [list(r) for r in rows]
    for index, row in enumerate(materialized[:12]):
        normalized = [str(v or "").strip().lower() for v in row]
        if any(alias in normalized for alias in BARCODE_ALIASES):
            return index, [str(v or "").strip() for v in row], materialized
    raise RosterError("No GotPhoto Barcode (1), Barcode, Copyright, or Code column was found")


def _extract(rows: list[list[object]]) -> RosterCodes:
    header_index, header, rows = _choose_header(rows)
    normalized = [h.lower() for h in header]
    alias = next(a for a in BARCODE_ALIASES if a in normalized)
    column_index = normalized.index(alias)
    values = [normalize_code(r[column_index] if column_index < len(r) else "")
              for r in rows[header_index + 1:]]
    codes = frozenset(v for v in values if v)
    if not codes:
        raise RosterError(f"The {header[column_index]!r} column contains no usable codes")
    return RosterCodes(codes, sum(bool(v) for v in values), header[column_index])


def load_roster_codes(path: Path) -> RosterCodes:
    path = Path(path)
    if not path.is_file():
        raise RosterError(f"Roster does not exist: {path}")
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    return _extract(list(csv.reader(handle)))
            except UnicodeDecodeError:
                continue
        raise RosterError("Roster CSV encoding is unsupported")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        errors: list[str] = []
        for sheet in workbook.worksheets:
            rows = [list(r) for r in sheet.iter_rows(values_only=True) if any(v is not None for v in r)]
            if not rows:
                continue
            try:
                return _extract(rows)
            except RosterError as exc:
                errors.append(f"{sheet.title}: {exc}")
        raise RosterError("No usable worksheet found. " + "; ".join(errors))
    raise RosterError("Roster must be CSV or XLSX")

